#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Local bilingual (Indonesian/English) PDF/PPTX/DOCX -> mT5 AQG -> quiz pipeline.

No paid API is used. For model-backed generation, pass a locally saved fine-tuned
mT5 directory through --model-dir. Without model weights, --demo uses an explicit
cloze-style fallback and labels each item's question_source as template_fallback.
"""
from __future__ import annotations

import argparse
import json
import math
import os
import random
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence

try:
    import fitz  # PyMuPDF
except Exception:
    fitz = None
try:
    from pptx import Presentation
except Exception:
    Presentation = None
try:
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except Exception:
    Document = None
    CT_Tbl = CT_P = Table = Paragraph = None
try:
    from langdetect import DetectorFactory, detect
    DetectorFactory.seed = 0
except Exception:
    detect = None

# mT5 is imported only when a model directory is actually supplied.

SUPER_DIGITS = str.maketrans({
    "⁰": "^0", "¹": "^1", "²": "^2", "³": "^3", "⁴": "^4",
    "⁵": "^5", "⁶": "^6", "⁷": "^7", "⁸": "^8", "⁹": "^9",
    "₀": "_0", "₁": "_1", "₂": "_2", "₃": "_3", "₄": "_4",
    "₅": "_5", "₆": "_6", "₇": "_7", "₈": "_8", "₉": "_9",
})
UNICODE_FRACTIONS = {
    "½": "1/2", "⅓": "1/3", "⅔": "2/3", "¼": "1/4",
    "¾": "3/4", "⅕": "1/5", "⅖": "2/5", "⅗": "3/5",
    "⅘": "4/5", "⅙": "1/6", "⅚": "5/6", "⅛": "1/8",
    "⅜": "3/8", "⅝": "5/8", "⅞": "7/8",
}
UNICODE_REPLACEMENTS = {
    "\u00ad": "", "\ufb00": "ff", "\ufb01": "fi", "\ufb02": "fl",
    "\ufb03": "ffi", "\ufb04": "ffl", "\ufb05": "st", "\ufb06": "st",
    "−": "-", "–": "-", "—": "-", "‑": "-", "﹣": "-",
    "×": " × ", "∙": " × ", "⋅": " × ", "·": " × ",
    "÷": " / ", "⁄": "/", "＝": "=", "≤": " <= ", "≥": " >= ",
    "≠": " != ", "≈": " approximately ", "π": "π", "ρ": "ρ", "μ": "μ",
    "α": "α", "β": "β", "γ": "γ", "θ": "θ", "λ": "λ", "Δ": "Δ",
    "Ω": "Ω", "ω": "ω", "√": "sqrt", "∞": "infinity",
}

FORMULA_PATTERNS = [
    re.compile(r"\$\$.*?\$\$", re.DOTALL),
    re.compile(r"\$(?!\$)[^$\n]{1,180}\$(?!\$)"),
    re.compile(r"\\\([^\n]{1,180}?\\\)"),
    re.compile(r"\\\[[\s\S]{1,240}?\\\]"),
    re.compile(r"\\(?:frac|dfrac|tfrac|sqrt|sum|int|prod|lim|partial|nabla|alpha|beta|gamma|theta|pi|sigma|mu|rho)\b(?:\s*\{[^}]*\})?"),
    re.compile(r"(?<!\w)[A-Za-zα-ωΑ-Ω][A-Za-z0-9_{}^]*\s*=\s*[^.;,\n]{1,100}"),
    re.compile(r"\b(?:d/dx|dy/dx|f'\s*\(\s*x\s*\)|∫|∑|Σ|sqrt\s*\()\s*[^.;,\n]{0,100}"),
    re.compile(r"(?<!\w)\d+(?:[.,]\d+)?\s*(?:\+|-|\*|/|×|÷|\^|<=|>=|!=)\s*\d+(?:[.,]\d+)?"),
]

DATE_PATTERNS = [
    re.compile(r"\b(?:19|20)\d{2}[-/.](?:0?[1-9]|1[0-2])[-/.](?:0?[1-9]|[12]\d|3[01])\b"),
    re.compile(r"\b(?:0?[1-9]|[12]\d|3[01])[-/.](?:0?[1-9]|1[0-2])[-/.](?:19|20)\d{2}\b"),
    re.compile(r"\b(?:19|20)\d{2}\b"),
]
NUMBER_UNIT_PATTERN = re.compile(
    r"(?<![\w^])[-+]?\d+(?:[.,]\d+)?\s*(?:%|persen|°\s?[CFK]|"
    r"kg/m(?:\^?\d+)?|g/mol|km/h|m/s(?:\^?\d+)?|cm(?:\^?\d+)?|mm(?:\^?\d+)?|km(?:\^?\d+)?|"
    r"m(?:\^?\d+)?|kg|mg|g|lb|N|J|W|kW|V|mV|A|mA|Ω|ohm|Pa|kPa|bar|"
    r"mol/L|mol|M|L|mL|cm³|m³|Hz|kHz|s|ms|min|menit|jam|hari|tahun|"
    r"rupiah|IDR|USD|\$|€|Rp\.?)(?!\w)?",
    re.IGNORECASE,
)
NUMBER_PLAIN_PATTERN = re.compile(r"(?<![\w^])[-+]?\d+(?:[.,]\d+)?(?![\w^])")

ID_STOPWORDS = set("yang dan di ke dari untuk dengan pada adalah merupakan atau dalam ini itu sebagai oleh karena dapat jika maka sebuah suatu data nilai rumus yaitu terhadap akan tidak dari juga pada secara".split())
EN_STOPWORDS = set("the and is are was were of to for with in on from as this that by can if then a an be value formula equation data into which what how when where".split())


def normalize_space(text: Any) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def normalize_math_text(text: Any) -> str:
    """Normalize common PDF math glyphs without discarding exponent/subscript meaning."""
    value = str(text or "").translate(SUPER_DIGITS)
    for old, new in UNICODE_FRACTIONS.items():
        value = value.replace(old, new)
    for old, new in UNICODE_REPLACEMENTS.items():
        value = value.replace(old, new)
    value = unicodedata.normalize("NFKC", value)
    value = value.replace("\u00a0", " ")
    # Fix common OCR/extraction artifacts while keeping Greek symbols readable.
    # Keep unit strings such as m/s and mol/L intact. A slash is padded only
    # when it behaves like a division operator between numeric/formula tokens.
    value = re.sub(r"\s*([=+*^])\s*", r" \1 ", value)
    value = re.sub(r"(?<=\d)\s*/\s*(?=[\dA-Za-z(])", " / ", value)
    value = re.sub(r"(?<=[A-Za-z0-9})])\s*/\s*(?=[A-Za-z0-9({])", lambda m: m.group(0) if False else " / ", value)
    # Undo spacing for common compound units after formula-operator normalization.
    value = re.sub(r"\b(m|km|cm|mm|kg|g|mol)\s*/\s*(s|h|L|mol|m)\s*\^\s*(\d+)\b", r"\1/\2^\3", value, flags=re.IGNORECASE)
    value = re.sub(r"\b(m|km|cm|mm|kg|g|mol)\s*/\s*(s|h|L|mol|m)\b", r"\1/\2", value, flags=re.IGNORECASE)
    value = re.sub(r"\b(cm|mm|km|m|kg|g|mol|L)\s*\^\s*(\d+)\b", r"\1^\2", value, flags=re.IGNORECASE)
    value = re.sub(r"\s*×\s*", " × ", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value


def detect_formula_spans(text: str) -> list[tuple[int, int, str]]:
    value = str(text or "")
    spans: list[tuple[int, int, str]] = []
    for pattern in FORMULA_PATTERNS:
        for match in pattern.finditer(value):
            raw = match.group(0).strip()
            if not raw:
                continue
            start, end = match.span()
            # Avoid marking a plain unit such as cm^2 as an equation.
            if re.fullmatch(r"[A-Za-z]{1,8}\s*\^\s*\d+", raw) and not any(c in raw for c in "=$\\"):
                continue
            spans.append((start, end, value[start:end]))
    spans.sort(key=lambda item: (item[0], -(item[1] - item[0])))
    merged: list[tuple[int, int, str]] = []
    for item in spans:
        start, end, raw = item
        if any(a <= start and b >= end for a, b, _ in merged):
            continue
        # Partial overlaps are consolidated so reverse insertions remain index-safe.
        overlap_idx = next((i for i, (a, b, _) in enumerate(merged) if start < b and end > a), None)
        if overlap_idx is not None:
            a, b, old = merged[overlap_idx]
            ns, ne = min(start, a), max(end, b)
            merged[overlap_idx] = (ns, ne, value[ns:ne])
            continue
        merged.append(item)
    merged.sort(key=lambda item: item[0])
    return merged[:30]


def answer_type(answer: str) -> str:
    value = normalize_space(answer)
    if not value:
        return "text"
    if any(p.fullmatch(value) for p in DATE_PATTERNS):
        return "date"
    # Number + unit should be classified before formula detection (e.g. 40 cm^2).
    if len(value.split()) <= 4 and re.fullmatch(
        r"[-+]?\d+(?:[.,]\d+)?\s*(?:%|persen|°\s?[CFK]|km/h|m/s(?:\^?\d+)?|"
        r"kg/m(?:\^?\d+)?|g/mol|cm(?:\^?\d+)?|mm(?:\^?\d+)?|km(?:\^?\d+)?|m(?:\^?\d+)?|kg|mg|g|"
        r"N|J|W|kW|V|mV|A|mA|Ω|ohm|Pa|kPa|bar|mol/L|mol|M|L|mL|Hz|kHz|s|ms|min|"
        r"menit|jam|hari|tahun|rupiah|IDR|USD|\$|€|Rp\.?|cm³|m³)?",
        value, re.IGNORECASE,
    ):
        return "number"
    if detect_formula_spans(value):
        return "formula"
    return "text"


def _formula_conflicts(start: int, end: int, answer_span: Optional[tuple[int, int]]) -> bool:
    if answer_span is None:
        return False
    a_start, a_end = answer_span
    return start < a_end and end > a_start


def mark_formulas(context: str, answer_span: Optional[tuple[int, int]] = None) -> str:
    """Wrap detected formula spans, excluding any span that overlaps the answer span."""
    text = str(context or "")
    spans = [s for s in detect_formula_spans(text) if not _formula_conflicts(s[0], s[1], answer_span)]
    for start, end, _ in sorted(spans, key=lambda x: x[0], reverse=True):
        text = text[:start] + "<formula>" + text[start:end] + "</formula>" + text[end:]
    return text


def mark_answer(context: str, answer: str, answer_start: int) -> str:
    """Mark an answer and preserve every character before and after it."""
    text, target = str(context or ""), str(answer or "")
    if not target:
        return text
    try:
        start = int(answer_start)
    except (TypeError, ValueError):
        start = -1
    if start < 0 or text[start:start + len(target)].casefold() != target.casefold():
        start = text.casefold().find(target.casefold())
    if start < 0:
        return text
    end = start + len(target)
    return text[:start] + "<answer>" + text[start:end] + "</answer>" + text[end:]


def mark_answer_and_formulas(context: str, answer: str, answer_start: int = -1) -> str:
    """Wrap answer plus formula spans using offsets from the unmodified context."""
    text, target = str(context or ""), str(answer or "")
    try:
        start = int(answer_start)
    except (TypeError, ValueError):
        start = -1
    if start < 0 or text[start:start + len(target)].casefold() != target.casefold():
        start = text.casefold().find(target.casefold()) if target else -1
    if start < 0:
        return mark_formulas(text, None)
    end = start + len(target)
    spans = detect_formula_spans(text)
    formula_spans = [s for s in spans if not _formula_conflicts(s[0], s[1], (start, end))]
    marked = mark_formulas(text, (start, end))
    # Earlier formula tags shift the answer position; later formulas do not.
    shift = sum(len("<formula>") + len("</formula>") for a, b, _ in formula_spans if b <= start)
    return mark_answer(marked, target, start + shift)


def convert_to_slide_style(context: str, answer: Optional[str] = None) -> str:
    """Convert prose into a compact slide-like topic + bullet format."""
    text = normalize_space(context)
    if not text:
        return text
    sentences = [x.strip(" -•\t") for x in re.split(r"(?<=[.!?])\s+|\s*[•▪]\s*", text) if x.strip(" -•\t")]
    if not sentences:
        sentences = [text]
    topic = sentences[0].rstrip(". ")
    bullets = sentences[1:] or sentences[:1]
    if answer and not any(str(answer).casefold() in s.casefold() for s in sentences):
        bullets.append(str(answer))
    return "Topik: " + topic + " | Poin Penting: " + " ".join("- " + x.rstrip(". ") + "." for x in bullets)


def detect_language(text: str) -> str:
    value = normalize_space(text)
    if not value:
        return "id"
    if detect is not None and len(value) >= 25:
        try:
            language = detect(value)
            if language == "en":
                return "en"
            if language == "id":
                return "id"
        except Exception:
            pass
    words = re.findall(r"[A-Za-zÀ-ÿ]+", value.casefold())
    id_score = sum(w in ID_STOPWORDS for w in words)
    en_score = sum(w in EN_STOPWORDS for w in words)
    if en_score > id_score:
        return "en"
    return "id"


def _word_count(text: str) -> int:
    return len(re.findall(r"\b\w+\b", str(text or ""), flags=re.UNICODE))


def _sentence_index(text: str, offset: int) -> int:
    return len(re.findall(r"[.!?](?:\s|$)", text[:max(0, offset)]))


def _clean_answer_span(span: str) -> str:
    return normalize_space(span).strip(" \t\r\n.,;:()[]{}")


class DocumentToContextPipeline:
    """Extract PDF/PPTX/DOCX content into contextual bilingual chunks."""

    def __init__(self, max_words: int = 220, min_slide_words: int = 25):
        self.max_words = max(60, int(max_words))
        self.min_slide_words = max(5, int(min_slide_words))

    def process(self, file_path: str | os.PathLike[str]) -> list[dict[str, Any]]:
        path = Path(file_path)
        if not path.exists() or not path.is_file():
            raise FileNotFoundError(f"File input tidak ditemukan: {path}")
        suffix = path.suffix.casefold()
        if suffix == ".pdf":
            records = self._extract_pdf(path)
        elif suffix == ".pptx":
            records = self._extract_pptx(path)
        elif suffix == ".docx":
            records = self._extract_docx(path)
        else:
            raise ValueError("Format tidak didukung. Gunakan file .pdf, .pptx, atau .docx.")
        chunks = self._chunk_records(records, path.name)
        for i, chunk in enumerate(chunks, start=1):
            chunk["chunk_id"] = i
            chunk["lang"] = detect_language(chunk["context"])
            chunk["word_count"] = _word_count(chunk["context"])
        return chunks

    def _extract_pdf(self, path: Path) -> list[dict[str, Any]]:
        if fitz is None:
            raise RuntimeError("PyMuPDF belum terpasang. Jalankan: pip install pymupdf")
        records: list[dict[str, Any]] = []
        with fitz.open(path) as pdf:
            page_lines: list[list[str]] = []
            for page in pdf:
                raw = page.get_text("text") or ""
                lines = [normalize_space(x) for x in raw.splitlines() if normalize_space(x)]
                page_lines.append(lines)
            # Remove repeated headers/footers found at the very top/bottom of multiple pages.
            counter: Counter[str] = Counter()
            for lines in page_lines:
                # Count a line at most once per page (first/last slices can overlap).
                edge_lines = set(line.casefold() for line in (lines[:2] + lines[-2:]) if len(line) <= 140)
                counter.update(edge_lines)
            threshold = max(2, math.ceil(len(page_lines) * 0.5))
            repeating = {line for line, count in counter.items() if count >= threshold} if len(page_lines) >= 2 else set()
            for page_index, (page, lines) in enumerate(zip(pdf, page_lines), start=1):
                cleaned = []
                for line_idx, line in enumerate(lines):
                    if line.casefold() in repeating and (line_idx < 2 or line_idx >= len(lines) - 2):
                        continue
                    if re.fullmatch(r"(?:halaman\s*)?\d+(?:\s*(?:/|dari)\s*\d+)?", line, re.IGNORECASE):
                        continue
                    if re.fullmatch(r"page\s+\d+(?:\s+of\s+\d+)?", line, re.IGNORECASE):
                        continue
                    cleaned.append(line)
                page_text = normalize_math_text("\n".join(cleaned))
                if not page_text:
                    continue
                # Capture bold spans from PDF text-dict blocks for keyword prioritization.
                bold_terms: list[str] = []
                try:
                    for block in page.get_text("dict").get("blocks", []):
                        for line in block.get("lines", []):
                            for span in line.get("spans", []):
                                font_name = str(span.get("font", "")).casefold()
                                if "bold" in font_name and normalize_space(span.get("text")):
                                    bold_terms.append(normalize_math_text(span["text"]))
                except Exception:
                    pass
                paragraphs = [normalize_space(p) for p in re.split(r"\n+", page_text) if normalize_space(p)]
                for paragraph in paragraphs:
                    records.append({"text": paragraph, "section": f"page_{page_index}", "bold_terms": bold_terms, "metadata": {"page": page_index}})
        return records

    @staticmethod
    def _table_to_sentences(table) -> list[str]:
        rows = [[normalize_space(DocumentToContextPipeline._xml_text(cell._tc) if hasattr(cell, "_tc") else cell.text) for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(row)]
        if not rows:
            return []
        headers = rows[0] if len(rows) > 1 else []
        content_rows = rows[1:] if headers else rows
        statements = []
        for row in content_rows:
            if headers and len(headers) == len(row):
                pairs = [f"{h}: {v}" for h, v in zip(headers, row) if h and v]
                if pairs:
                    statements.append("Data pada tabel menunjukkan " + ", ".join(pairs) + ".")
            else:
                statements.append("Data pada tabel: " + "; ".join(x for x in row if x) + ".")
        return statements

    @staticmethod
    def _get_speaker_notes(slide) -> str:
        try:
            if not slide.has_notes_slide:
                return ""
            frame = slide.notes_slide.notes_text_frame
            if frame is None:
                return ""
            lines = [normalize_space(p.text) for p in frame.paragraphs if normalize_space(p.text)]
            lines = [x for x in lines if x.casefold() not in {"slide number", "click to edit master text styles", "click to add notes"}]
            return normalize_space(" ".join(lines))
        except Exception:
            return ""

    def _extract_pptx(self, path: Path) -> list[dict[str, Any]]:
        if Presentation is None:
            raise RuntimeError("python-pptx belum terpasang. Jalankan: pip install python-pptx")
        presentation = Presentation(str(path))
        slides: list[dict[str, Any]] = []
        for index, slide in enumerate(presentation.slides, start=1):
            title_shape = slide.shapes.title if slide.shapes.title is not None else None
            title = normalize_math_text(title_shape.text) if title_shape is not None else ""
            pieces: list[str] = []
            bold_terms: list[str] = []
            for shape in slide.shapes:
                if title_shape is not None and shape == title_shape:
                    continue
                if getattr(shape, "has_text_frame", False):
                    for para in shape.text_frame.paragraphs:
                        text = normalize_math_text(para.text)
                        if text:
                            pieces.append(text)
                        for run in para.runs:
                            if run.font.bold and normalize_space(run.text):
                                bold_terms.append(normalize_math_text(run.text))
                if getattr(shape, "has_table", False):
                    pieces.extend(self._table_to_sentences(shape.table))
            notes = self._get_speaker_notes(slide)
            body_parts = []
            if title:
                body_parts.append(f"Pada materi {title}, dijelaskan bahwa:")
            if pieces:
                body_parts.append(". Selain itu, ".join(p.strip(" .") for p in pieces) + ".")
            if notes:
                body_parts.append("Catatan tambahan: " + notes)
            content = normalize_math_text(" ".join(body_parts))
            if content:
                slides.append({
                    "title": title or f"Slide {index}", "text": content,
                    "bold_terms": bold_terms, "slide_start": index, "slide_end": index,
                })
        # Merge only short slides when titles explicitly repeat or indicate continuation.
        merged: list[dict[str, Any]] = []
        for slide in slides:
            if merged and _word_count(merged[-1]["text"]) < self.min_slide_words and self._titles_continue(merged[-1]["title"], slide["title"]):
                previous = merged[-1]
                previous["text"] = normalize_space(previous["text"] + " Selain itu, " + slide["text"])
                previous["bold_terms"].extend(slide["bold_terms"])
                previous["slide_end"] = slide["slide_end"]
                if previous["title"].casefold().startswith("lanjutan"):
                    previous["title"] = slide["title"]
            else:
                merged.append(slide.copy())
        return [{
            "text": s["text"], "section": s["title"], "bold_terms": s["bold_terms"],
            "metadata": {"slide_start": s["slide_start"], "slide_end": s["slide_end"]},
        } for s in merged]

    @staticmethod
    def _titles_continue(title_a: str, title_b: str) -> bool:
        a = re.sub(r"[^a-z0-9 ]", "", normalize_space(title_a).casefold())
        b = re.sub(r"[^a-z0-9 ]", "", normalize_space(title_b).casefold())
        if not a or not b:
            return False
        if a == b or a in b or b in a:
            return True
        continuation_words = ("lanjutan", "continued", "part 2", "bagian 2", "(2)", "- 2")
        if any(x in title_b.casefold() for x in continuation_words):
            base = re.sub(r"\b(lanjutan|continued|part 2|bagian 2)\b", "", b).strip()
            return not base or bool(base and (base in a or a in base))
        # A continuation with same first two meaningful words is a conservative match.
        aw = [w for w in a.split() if len(w) > 3]
        bw = [w for w in b.split() if len(w) > 3]
        return len(set(aw[:2]) & set(bw[:2])) >= 2 if len(aw) >= 2 and len(bw) >= 2 else False

    @staticmethod
    def _xml_text(node) -> str:
        """Read WordprocessingML text while converting OMML equations to readable text."""
        tag = node.tag.split("}")[-1]
        if tag in {"oMath", "oMathPara"}:
            return DocumentToContextPipeline._omml_to_text(node)
        if tag == "t":
            return node.text or ""
        if tag in {"tab"}:
            return "\t"
        if tag in {"br", "cr"}:
            return " "
        # Do not recurse into deleted text or field instruction content.
        if tag in {"del", "instrText", "delText"}:
            return ""
        return "".join(DocumentToContextPipeline._xml_text(child) for child in list(node))

    @staticmethod
    def _omml_to_text(node) -> str:
        tag = node.tag.split("}")[-1]
        children = list(node)
        def part(local: str) -> str:
            found = next((child for child in children if child.tag.split("}")[-1] == local), None)
            return DocumentToContextPipeline._omml_to_text(found) if found is not None else ""
        if tag == "t":
            return node.text or ""
        if tag == "f":
            num, den = part("num"), part("den")
            return f"({num})/({den})" if num and den else "fraction"
        if tag in {"sSup", "sSub", "sSubSup"}:
            base, sup, sub = part("e"), part("sup"), part("sub")
            if tag == "sSup":
                return f"{base}^{sup}"
            if tag == "sSub":
                return f"{base}_{sub}"
            return f"{base}_{sub}^{sup}"
        if tag == "rad":
            degree, value = part("deg"), part("e")
            return f"root[{degree}]({value})" if degree else f"sqrt({value})"
        if tag == "nary":
            nary_props = next((child for child in children if child.tag.split("}")[-1] == "naryPr"), None)
            operator = "∑"
            if nary_props is not None:
                for child in nary_props.iter():
                    local = child.tag.split("}")[-1]
                    if local == "chr":
                        raw_val = next((v for k, v in child.attrib.items() if k.endswith("}val") or k == "val"), "")
                        if raw_val in {"∫", "\\int"}:
                            operator = "∫"
                            break
            lower, upper, value = part("sub"), part("sup"), part("e")
            limits = "".join(x for x in [f"_{{{lower}}}" if lower else "", f"^{{{upper}}}" if upper else ""])
            return f"{operator}{limits} {value}".strip()
        if tag == "d":
            return "(" + "".join(DocumentToContextPipeline._omml_to_text(x) for x in children if x.tag.split("}")[-1] not in {"dPr"}) + ")"
        if tag == "func":
            name, arg = part("fName"), part("e")
            return f"{name}({arg})"
        return "".join(DocumentToContextPipeline._omml_to_text(child) for child in children)

    def _extract_docx(self, path: Path) -> list[dict[str, Any]]:
        if Document is None:
            raise RuntimeError("python-docx belum terpasang. Jalankan: pip install python-docx")
        doc = Document(str(path))
        records: list[dict[str, Any]] = []
        body = doc.element.body
        paragraph_index = 0
        table_index = 0
        for child in body.iterchildren():
            if CT_P is not None and isinstance(child, CT_P):
                para = Paragraph(child, doc)
                text = normalize_math_text(self._xml_text(child))
                if text:
                    bold_terms = [normalize_math_text(run.text) for run in para.runs if run.bold and normalize_space(run.text)]
                    paragraph_index += 1
                    records.append({"text": text, "section": f"paragraph_{paragraph_index}", "bold_terms": bold_terms, "metadata": {"paragraph": paragraph_index}})
            elif CT_Tbl is not None and isinstance(child, CT_Tbl):
                table = Table(child, doc)
                sentences = self._table_to_sentences(table)
                table_index += 1
                for row_index, sentence in enumerate(sentences, start=1):
                    records.append({"text": normalize_math_text(sentence), "section": f"table_{table_index}", "bold_terms": [], "metadata": {"table": table_index, "row": row_index}})
        return records

    def _chunk_records(self, records: Sequence[dict[str, Any]], source_name: str) -> list[dict[str, Any]]:
        chunks: list[dict[str, Any]] = []
        current_text: list[str] = []
        current_bold: list[str] = []
        current_sections: list[str] = []
        current_meta: list[dict[str, Any]] = []
        count = 0

        def flush() -> None:
            nonlocal current_text, current_bold, current_sections, current_meta, count
            combined = normalize_space(" ".join(current_text))
            if combined:
                chunks.append({
                    "context": combined,
                    "source_file": source_name,
                    "section": " | ".join(dict.fromkeys(current_sections)),
                    "bold_terms": list(dict.fromkeys(current_bold)),
                    "metadata": current_meta.copy(),
                })
            current_text, current_bold, current_sections, current_meta, count = [], [], [], [], 0

        for record in records:
            text = normalize_math_text(record.get("text", ""))
            if not text:
                continue
            word_n = _word_count(text)
            if word_n > self.max_words:
                flush()
                words = text.split()
                for begin in range(0, len(words), self.max_words):
                    segment = " ".join(words[begin:begin + self.max_words])
                    chunks.append({
                        "context": segment, "source_file": source_name,
                        "section": str(record.get("section", "")),
                        "bold_terms": list(record.get("bold_terms", [])),
                        "metadata": [record.get("metadata", {})],
                    })
                continue
            if current_text and count + word_n > self.max_words:
                flush()
            current_text.append(text)
            current_bold.extend(record.get("bold_terms", []))
            current_sections.append(str(record.get("section", "")))
            current_meta.append(record.get("metadata", {}))
            count += word_n
        flush()
        return chunks


class CandidateAnswerExtractor:
    """Heuristic candidate answer detection for formula, number/date, and key concepts."""

    DEFINITION_PATTERNS = {
        "id": re.compile(r"(?P<left>[^.!?;\n]{2,90}?)\s+(?P<cue>adalah|merupakan|disebut|didefinisikan\s+sebagai)\s+(?P<right>[^.!?;\n]{2,120})", re.IGNORECASE),
        "en": re.compile(r"(?P<left>[^.!?;\n]{2,90}?)\s+(?P<cue>is\s+defined\s+as|refers\s+to|is\s+called|known\s+as|is)\s+(?P<right>[^.!?;\n]{2,120})", re.IGNORECASE),
    }
    CAPITALIZED_PHRASE = re.compile(r"\b(?:[A-Z][A-Za-zÀ-ÿ0-9_-]+)(?:\s+(?:[A-Z][A-Za-zÀ-ÿ0-9_-]+|[a-z]{1,3})){0,3}\b")

    def extract(self, context: str, lang: Optional[str] = None, bold_terms: Optional[Sequence[str]] = None, max_candidates: int = 5) -> list[dict[str, Any]]:
        text = normalize_math_text(context)
        language = lang if lang in {"id", "en"} else detect_language(text)
        candidates: list[dict[str, Any]] = []

        def add(start: int, end: int, raw: str, kind: str, score: float, source: str) -> None:
            answer = _clean_answer_span(raw)
            if not answer or len(answer) > 180 or _word_count(answer) > 18:
                return
            if len(answer) == 1 and answer.casefold() in ID_STOPWORDS | EN_STOPWORDS:
                return
            inferred_kind = answer_type(answer)
            if kind == "text" and inferred_kind != "text":
                kind = inferred_kind
                score = min(score, 6.0)
            candidates.append({
                "answer": answer, "answer_type": kind, "start": max(0, int(start)),
                "end": max(0, int(end)), "score": float(score), "source": source,
                "sentence_index": _sentence_index(text, start),
            })

        formula_spans = detect_formula_spans(text)
        for start, end, value in formula_spans:
            add(start, end, value, "formula", 9.5, "formula_span")

        date_spans: list[tuple[int, int]] = []
        for pattern in DATE_PATTERNS:
            for match in pattern.finditer(text):
                date_spans.append(match.span())
        for start, end in sorted(set(date_spans)):
            add(start, end, text[start:end], "date", 8.4, "date_pattern")

        numeric_spans: list[tuple[int, int]] = []
        for match in NUMBER_UNIT_PATTERN.finditer(text):
            span = match.span()
            if any(a <= span[0] and b >= span[1] for a, b, _ in formula_spans):
                continue
            if any(a <= span[0] and b >= span[1] for a, b in date_spans):
                continue
            numeric_spans.append(span)
            add(span[0], span[1], match.group(0), "number", 7.7, "number_with_unit")
        # Plain number candidates are lower priority and avoid values already captured with units/dates/formulas.
        for match in NUMBER_PLAIN_PATTERN.finditer(text):
            span = match.span()
            if any(a <= span[0] and b >= span[1] for a, b in numeric_spans + date_spans):
                continue
            if any(a <= span[0] and b >= span[1] for a, b, _ in formula_spans):
                continue
            add(span[0], span[1], match.group(0), "number", 4.8, "plain_number")

        definition_pattern = self.DEFINITION_PATTERNS[language]
        for match in definition_pattern.finditer(text):
            left = _clean_answer_span(match.group("left"))
            right = _clean_answer_span(match.group("right"))
            # Keep concise concepts, not sentence-sized definitions.
            if 1 <= _word_count(left) <= 8:
                start = match.start("left") + len(match.group("left")) - len(match.group("left").lstrip())
                add(start, start + len(left), left, "text", 6.8, "definition_subject")
            right_core = re.split(r"\s+(?:yang|yang\s+berfungsi|which|that|with|untuk|untuk\s+melakukan|sehingga|seperti)\s+", right, maxsplit=1, flags=re.IGNORECASE)[0]
            if 2 <= _word_count(right_core) <= 10:
                start = match.start("right")
                add(start, start + len(right_core), right_core, "text", 6.5, "definition_object")

        bold_lookup = {normalize_space(term).casefold(): term for term in (bold_terms or []) if normalize_space(term)}
        for term in bold_terms or []:
            term = normalize_space(term)
            if not term:
                continue
            offset = text.casefold().find(term.casefold())
            if offset >= 0:
                add(offset, offset + len(term), term, "text", 7.2, "bold_text")

        for match in self.CAPITALIZED_PHRASE.finditer(text):
            phrase = _clean_answer_span(match.group(0))
            words = phrase.split()
            if _word_count(phrase) >= 2 and words[-1].strip(".,;:").casefold() not in {"v", "a", "x", "y", "z"} and phrase.casefold() not in bold_lookup:
                add(match.start(), match.end(), phrase, "text", 4.5, "capitalized_phrase")

        # General technical terms in quotes or after label markers.
        for match in re.finditer(r"(?:istilah|konsep|variabel|parameter|concept|term|variable|parameter)\s*[:=]\s*([^.;,\n]{2,70})", text, re.IGNORECASE):
            value = _clean_answer_span(match.group(1))
            if 1 <= _word_count(value) <= 8:
                add(match.start(1), match.end(1), value, "text", 5.8, "labeled_concept")

        # Deduplicate normalized answer/type pairs, keeping the highest score.
        unique: dict[tuple[str, str], dict[str, Any]] = {}
        for candidate in candidates:
            key = (re.sub(r"\W+", " ", candidate["answer"]).strip().casefold(), candidate["answer_type"])
            existing = unique.get(key)
            if existing is None or candidate["score"] > existing["score"]:
                unique[key] = candidate
        pool = list(unique.values())
        if not pool:
            return []
        # Select across the context instead of taking many candidates from one sentence.
        selected: list[dict[str, Any]] = []
        pool.sort(key=lambda c: (-c["score"], c["start"]))
        limit = max(1, min(int(max_candidates), 5))
        while pool and len(selected) < limit:
            best_idx, best_value = 0, -1e9
            for idx, candidate in enumerate(pool):
                min_gap = min((abs(candidate["start"] - old["start"]) for old in selected), default=len(text))
                same_sentence = any(candidate["sentence_index"] == old["sentence_index"] for old in selected)
                distance_bonus = min(min_gap / max(1, len(text)), 0.35) * 2.0
                sentence_penalty = 1.5 if same_sentence else 0.0
                value = candidate["score"] + distance_bonus - sentence_penalty
                if value > best_value:
                    best_idx, best_value = idx, value
            selected.append(pool.pop(best_idx))
        selected.sort(key=lambda c: c["start"])
        return selected


def _parse_numeric_answer(answer: str) -> Optional[tuple[float, str, str]]:
    value = normalize_space(answer)
    match = re.fullmatch(r"\s*([-+]?\d+(?:[.,]\d+)?)\s*(.*?)\s*", value)
    if not match:
        return None
    try:
        number = float(match.group(1).replace(",", "."))
    except ValueError:
        return None
    unit = match.group(2)
    decimal_sep = "," if "," in match.group(1) else "."
    return number, unit, decimal_sep


def _format_number(value: float, decimal_sep: str = ".") -> str:
    if not math.isfinite(value):
        return "0"
    if abs(value - round(value)) < 1e-9:
        return str(int(round(value)))
    output = f"{value:.3f}".rstrip("0").rstrip(".")
    return output.replace(".", decimal_sep)


def _all_answer_records(all_doc_answers: Iterable[Any]) -> list[dict[str, Any]]:
    rows = []
    for item in all_doc_answers or []:
        if isinstance(item, str):
            rows.append({"answer": item, "answer_type": answer_type(item)})
        elif isinstance(item, dict) and item.get("answer"):
            rows.append({"answer": str(item["answer"]), "answer_type": item.get("answer_type", answer_type(str(item["answer"])))})
    return rows


def generate_distractors(answer: str, answer_type_name: str, context: str, all_doc_answers: Iterable[Any]) -> list[str]:
    """Generate three candidate distractors with same broad answer type and unit."""
    target = normalize_space(answer)
    kind = answer_type_name or answer_type(target)
    pool: list[str] = []
    if kind == "formula":
        formula = target.replace("<formula>", "").replace("</formula>", "")
        operator_positions = [(i, ch) for i, ch in enumerate(formula) if ch in "+-×*/÷"]
        operator_alternatives = {
            "+": ["-", "×", "/"], "-": ["+", "×", "/"],
            "×": ["+", "-", "/"], "*": ["+", "-", "/"],
            "/": ["×", "+", "-"], "÷": ["×", "+", "-"],
        }
        for position, operator in operator_positions:
            for replacement in operator_alternatives[operator]:
                altered = formula[:position] + replacement + formula[position + 1:]
                if altered != formula:
                    pool.append(altered)
                if len(pool) >= 9:
                    break
            if len(pool) >= 9:
                break
        if len(pool) < 3:
            variables = list(dict.fromkeys(re.findall(r"(?<![A-Za-z])(?:[A-Za-zα-ωΑ-Ω][A-Za-z0-9_]*)(?![A-Za-z])", formula)))
            for i in range(len(variables)):
                for j in range(i + 1, len(variables)):
                    a, b = variables[i], variables[j]
                    swapped = re.sub(r"\bTEMPVAR\b", a, formula.replace(a, "TEMPVAR", 1).replace(b, a, 1), count=1).replace("TEMPVAR", b, 1)
                    if swapped != formula:
                        pool.append(swapped)
        pool.extend([r["answer"] for r in _all_answer_records(all_doc_answers) if r["answer_type"] == "formula" and r["answer"] != target])
    elif kind == "number":
        parsed = _parse_numeric_answer(target)
        if parsed:
            number, unit, sep = parsed
            for number_variant in (number * 0.5, number * 2.0, number + 10 if number >= 0 else number - 10, number + 5):
                candidate = _format_number(number_variant, sep) + (" " + unit if unit else "")
                pool.append(candidate)
        pool.extend([r["answer"] for r in _all_answer_records(all_doc_answers) if r["answer_type"] == "number" and r["answer"] != target])
    elif kind == "date":
        full_iso = re.fullmatch(r"((?:19|20)\d{2})[-/.](\d{1,2})[-/.](\d{1,2})", target)
        full_dmy = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})[-/.]((?:19|20)\d{2})", target)
        year_only = re.fullmatch(r"((?:19|20)\d{2})", target)
        if year_only:
            year = int(year_only.group(1))
            pool.extend(str(year + delta) for delta in (-1, 1, -2, 2, 5, -5, 10, -10) if 1800 <= year + delta <= 2200)
        elif full_iso:
            y, m, d = map(int, full_iso.groups())
            for delta in (-1, 1, -2, 2, 5):
                new_y = y + delta
                try:
                    pool.append(date(new_y, m, d).strftime("%Y-%m-%d"))
                except ValueError:
                    pool.append(date(new_y, m, min(d, 28)).strftime("%Y-%m-%d"))
        elif full_dmy:
            d, m, y = map(int, full_dmy.groups())
            for delta in (-1, 1, -2, 2, 5):
                new_y = y + delta
                try:
                    pool.append(date(new_y, m, d).strftime("%d/%m/%Y"))
                except ValueError:
                    pool.append(date(new_y, m, min(d, 28)).strftime("%d/%m/%Y"))
        pool.extend([r["answer"] for r in _all_answer_records(all_doc_answers) if r["answer_type"] == "date" and r["answer"] != target])
    else:
        records = [r for r in _all_answer_records(all_doc_answers) if r["answer_type"] == "text" and r["answer"].casefold() != target.casefold()]
        target_words = _word_count(target)
        records.sort(key=lambda r: (abs(_word_count(r["answer"]) - target_words), abs(len(r["answer"]) - len(target))))
        pool.extend(r["answer"] for r in records)

    result: list[str] = []
    target_key = normalize_space(target).casefold()
    for candidate in pool:
        value = normalize_space(candidate)
        if not value or value.casefold() == target_key or value.casefold() in {x.casefold() for x in result}:
            continue
        if kind == "number":
            p = _parse_numeric_answer(value)
            t = _parse_numeric_answer(target)
            if p and t and p[1] != t[1]:
                continue  # numeric distractors must preserve the exact unit string
        if kind == "text" and answer_type(value) != "text":
            continue
        result.append(value)
        if len(result) == 3:
            break
    return result


def answer_type_alignment(question: str, answer: str, lang: str) -> Optional[bool]:
    """Soft alignment diagnostic: False means clearly mismatched; None means uncertain."""
    q = normalize_space(question).casefold()
    kind = answer_type(answer)
    lang = lang if lang in {"id", "en"} else detect_language(question)
    if kind == "date":
        cues = ("kapan", "tahun berapa", "tanggal berapa", "pada tahun") if lang == "id" else ("when", "what year", "which year", "what date", "on what date")
        contrad = ("berapa volt", "berapa meter", "rumus apa") if lang == "id" else ("how many volts", "how many meters", "which formula")
        if any(c in q for c in contrad):
            return False
        return True if any(c in q for c in cues) else None
    if kind == "number":
        cues = ("berapa", "berapa banyak", "berapa nilai", "berapa besar", "berapa hasil", "nilai yang") if lang == "id" else ("how many", "how much", "what value", "what is the value", "how large", "how fast", "how long")
        contrad = ("tahun berapa", "tanggal berapa", "rumus apa", "persamaan apa") if lang == "id" else ("what year", "what date", "which formula", "what equation")
        if any(c in q for c in contrad):
            return False
        return True if any(c in q for c in cues) else None
    if kind == "formula":
        cues = ("rumus", "persamaan", "formula", "bentuk persamaan") if lang == "id" else ("formula", "equation", "expression")
        contrad = ("berapa", "tahun berapa", "tanggal berapa") if lang == "id" else ("how many", "how much", "what year", "what date")
        if any(c in q for c in contrad):
            return False
        return True if any(c in q for c in cues) else None
    return None


class QuizGenerator:
    """Candidate-driven QG, quality filtering, distractors, and JSON serialization."""

    def __init__(self, model_dir: Optional[str] = None, batch_size: int = 4, seed: int = 42, load_model: bool = True):
        self.model_dir = str(model_dir) if model_dir else None
        self.batch_size = max(1, int(batch_size))
        self.seed = int(seed)
        self.extractor = CandidateAnswerExtractor()
        self.tokenizer = None
        self.model = None
        self.device = "cpu"
        if load_model and self.model_dir:
            self._load_model(self.model_dir)

    def _load_model(self, model_dir: str) -> None:
        path = Path(model_dir)
        if not path.exists():
            raise FileNotFoundError(f"Direktori model tidak ditemukan: {model_dir}")
        try:
            import torch
            from transformers import AutoModelForSeq2SeqLM, AutoTokenizer
        except Exception as exc:
            raise RuntimeError("Model inference memerlukan torch dan transformers. Instal dependensi sesuai README.") from exc
        self.device = "cuda:0" if torch.cuda.is_available() else "cpu"
        self.tokenizer = AutoTokenizer.from_pretrained(model_dir, local_files_only=True)
        self.model = AutoModelForSeq2SeqLM.from_pretrained(model_dir, local_files_only=True)
        self.model.to(self.device)
        self.model.eval()

    @property
    def uses_model(self) -> bool:
        return self.model is not None and self.tokenizer is not None

    def _prompt(self, context: str, candidate: dict[str, Any], lang: str) -> str:
        answer = str(candidate["answer"])
        try:
            answer_start = int(candidate.get("start", -1))
        except (TypeError, ValueError):
            answer_start = -1
        if answer_start < 0 or context[answer_start:answer_start + len(answer)].casefold() != answer.casefold():
            answer_start = context.casefold().find(answer.casefold()) if answer else -1
        has_formula = bool(detect_formula_spans(context))

        def build(window: str, local_start: int) -> str:
            marked = mark_answer_and_formulas(window, answer, local_start)
            return (
                f"generate question {lang}: answer_type: {candidate['answer_type']}: "
                f"formula_context: {'yes' if has_formula else 'no'}: "
                f"answer: {answer} context: {marked}"
            )

        # Mirror the training notebook's answer-centred truncation to keep the target span visible.
        if self.tokenizer is None or answer_start < 0:
            return build(context, answer_start)
        left, right = 0, len(context)
        prompt = build(context, answer_start)
        for _ in range(10):
            try:
                token_count = len(self.tokenizer(prompt, truncation=False)["input_ids"])
            except Exception:
                return prompt
            if token_count <= 512:
                return prompt
            half_width = max(120, int((right - left) * 0.32))
            left = max(0, answer_start - half_width)
            right = min(len(context), answer_start + len(answer) + half_width)
            window = context[left:right]
            local_start = answer_start - left
            if window[local_start:local_start + len(answer)].casefold() != answer.casefold():
                local_start = window.casefold().find(answer.casefold())
            if local_start < 0:
                return prompt
            prompt = build(window, local_start)
        return prompt

    def _fallback_question(self, context: str, candidate: dict[str, Any], lang: str) -> str:
        answer = candidate["answer"]
        start = int(candidate.get("start", -1))
        if start < 0 or context[start:start + len(answer)].casefold() != answer.casefold():
            start = context.casefold().find(answer.casefold())
        if start >= 0:
            cloze = context[:start] + "_____" + context[start + len(answer):]
            cloze = normalize_space(cloze).strip(" .")
        else:
            cloze = context
        if lang == "en":
            return "Complete the statement based on the material: " + cloze
        return "Lengkapi pernyataan berdasarkan materi: " + cloze

    def _generate_questions(self, pairs: list[dict[str, Any]]) -> list[str]:
        if not self.uses_model:
            return [self._fallback_question(p["context"], p["candidate"], p["lang"]) for p in pairs]
        outputs: list[str] = []
        import torch
        for begin in range(0, len(pairs), self.batch_size):
            batch = pairs[begin:begin + self.batch_size]
            prompts = [self._prompt(p["context"], p["candidate"], p["lang"]) for p in batch]
            encoded = self.tokenizer(prompts, padding=True, truncation=True, max_length=512, return_tensors="pt")
            encoded = {key: value.to(self.device) for key, value in encoded.items()}
            with torch.inference_mode():
                generated = self.model.generate(
                    **encoded, max_new_tokens=64, num_beams=4, do_sample=False,
                    no_repeat_ngram_size=2, early_stopping=True,
                )
            decoded = self.tokenizer.batch_decode(generated, skip_special_tokens=True)
            outputs.extend(normalize_space(x) for x in decoded)
        return outputs

    @staticmethod
    def _leaks_answer(question: str, answer: str) -> bool:
        q = normalize_space(question).casefold()
        a = normalize_space(answer).casefold()
        if not a:
            return False
        kind = answer_type(answer)
        if kind == "formula":
            # Treat equivalent whitespace in formulas as a leak too (e.g. v=s/t vs v = s / t).
            q_compact = re.sub(r"\s+", "", q)
            a_compact = re.sub(r"\s+", "", a)
            return bool(a_compact and a_compact in q_compact)
        if len(a) <= 2:
            return bool(re.search(rf"(?<!\w){re.escape(a)}(?!\w)", q))
        return a in q

    def generate_from_chunks(self, chunks: Sequence[dict[str, Any]], max_candidates_per_chunk: int = 5) -> dict[str, Any]:
        candidate_groups: list[tuple[dict[str, Any], list[dict[str, Any]]]] = []
        all_candidates: list[dict[str, Any]] = []
        for chunk_index, chunk in enumerate(chunks, start=1):
            context = normalize_math_text(chunk.get("context", ""))
            if not context:
                continue
            lang = chunk.get("lang") if chunk.get("lang") in {"id", "en"} else detect_language(context)
            candidates = self.extractor.extract(
                context, lang=lang, bold_terms=chunk.get("bold_terms", []),
                max_candidates=max_candidates_per_chunk,
            )
            for candidate in candidates:
                candidate.update({"chunk_index": chunk_index, "context": context, "lang": lang})
                all_candidates.append(candidate)
            candidate_groups.append((chunk, candidates))
        # Build flattened input pairs so actual model inference is batched across a whole document.
        pairs = []
        for chunk, candidates in candidate_groups:
            context = normalize_math_text(chunk.get("context", ""))
            lang = chunk.get("lang") if chunk.get("lang") in {"id", "en"} else detect_language(context)
            for candidate in candidates:
                pairs.append({"context": context, "candidate": candidate, "lang": lang, "chunk": chunk})
        questions = self._generate_questions(pairs)
        quizzes: list[dict[str, Any]] = []
        seen_questions: set[tuple[str, str]] = set()
        for index, (pair, question) in enumerate(zip(pairs, questions), start=1):
            candidate = pair["candidate"]
            context = pair["context"]
            lang = pair["lang"]
            answer = candidate["answer"]
            question = normalize_space(question).strip(" \t\r\n-•")
            if not question or len(question) < 8:
                continue
            if self._leaks_answer(question, answer):
                continue
            alignment = answer_type_alignment(question, answer, lang)
            if alignment is False:
                continue
            dedupe_key = (question.casefold(), answer.casefold())
            if dedupe_key in seen_questions:
                continue
            seen_questions.add(dedupe_key)
            other_answers = [x for x in all_candidates if x["answer"].casefold() != answer.casefold()]
            distractors = generate_distractors(answer, candidate["answer_type"], context, other_answers)
            qtype = "multiple_choice" if len(distractors) >= 3 else "essay"
            options = None
            if qtype == "multiple_choice":
                option_rows = [{"text": answer, "is_correct": True}] + [{"text": item, "is_correct": False} for item in distractors]
                random.Random(self.seed + index).shuffle(option_rows)
                options = [{"label": chr(65 + j), **row} for j, row in enumerate(option_rows)]
            quizzes.append({
                "id": f"Q{len(quizzes) + 1:03d}",
                "lang": lang,
                "question": question,
                "question_type": qtype,
                "answer": answer,
                "answer_type": candidate["answer_type"],
                "options": options,
                "distractors": distractors if qtype == "multiple_choice" else [],
                "alignment_check": alignment,
                "candidate_score": round(float(candidate["score"]), 3),
                "candidate_source": candidate["source"],
                "question_source": "mt5" if self.uses_model else "template_fallback",
                "source_file": pair["chunk"].get("source_file"),
                "section": pair["chunk"].get("section"),
                "chunk_index": candidate["chunk_index"],
            })
        return {
            "schema_version": "1.0",
            "model": self.model_dir if self.uses_model else None,
            "generation_mode": "mT5-base" if self.uses_model else "template_fallback_no_model_weights",
            "input_chunks": len(candidate_groups),
            "candidate_count": len(all_candidates),
            "quiz_count": len(quizzes),
            "quizzes": quizzes,
        }

    def generate_from_file(self, file_path: str | os.PathLike[str], max_candidates_per_chunk: int = 5) -> dict[str, Any]:
        chunks = DocumentToContextPipeline().process(file_path)
        result = self.generate_from_chunks(chunks, max_candidates_per_chunk=max_candidates_per_chunk)
        result["input_file"] = str(file_path)
        result["document_chunks"] = chunks
        return result


def build_demo_chunks() -> list[dict[str, Any]]:
    """Small, self-contained Indonesian and English slide-style formula examples."""
    id_context = (
        "Topik: Gerak Lurus | Poin Penting: - Kecepatan rata-rata dihitung dengan rumus v = s/t. "
        "- Sebuah benda menempuh jarak 120 m dalam waktu 10 s, sehingga kecepatan rata-ratanya 12 m/s. "
        "- Kecepatan menunjukkan perubahan posisi terhadap waktu."
    )
    en_context = (
        "Topic: Linear Motion | Key Points: - Average speed is calculated using v = s/t. "
        "- An object travels 150 m in 5 s, so its average speed is 30 m/s. "
        "- Speed describes the rate of change of distance over time."
    )
    return [
        {"context": normalize_math_text(id_context), "lang": "id", "source_file": "sample_ppt_id", "section": "Gerak Lurus", "bold_terms": ["Gerak Lurus"]},
        {"context": normalize_math_text(en_context), "lang": "en", "source_file": "sample_ppt_en", "section": "Linear Motion", "bold_terms": ["Linear Motion"]},
    ]


def _json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="PDF/PPTX/DOCX to bilingual AQG quiz JSON")
    parser.add_argument("--input", help="Input .pdf, .pptx, or .docx")
    parser.add_argument("--model-dir", default=os.environ.get("AQG_MODEL_DIR"), help="Local fine-tuned mT5 model directory")
    parser.add_argument("--output", help="Output JSON path; default prints to stdout")
    parser.add_argument("--batch-size", type=int, default=4)
    parser.add_argument("--max-candidates", type=int, default=5)
    parser.add_argument("--demo", action="store_true", help="Run bilingual short-slide examples")
    parser.add_argument("--no-model", action="store_true", help="Use fallback questions even if --model-dir is set")
    args = parser.parse_args(argv)

    if not args.demo and not args.input:
        parser.print_help()
        return 2
    if args.input and args.demo:
        parser.error("Gunakan salah satu: --input atau --demo")
    use_model = bool(args.model_dir) and not args.no_model
    if use_model:
        generator = QuizGenerator(model_dir=args.model_dir, batch_size=args.batch_size, load_model=True)
    else:
        generator = QuizGenerator(batch_size=args.batch_size, load_model=False)
        print("PERINGATAN: bobot mT5 tidak dimuat; pertanyaan dibuat dengan fallback cloze, bukan hasil inferensi model.", file=sys.stderr)

    if args.demo:
        result = generator.generate_from_chunks(build_demo_chunks(), max_candidates_per_chunk=args.max_candidates)
        result["demo_note"] = "Untuk output mT5, jalankan ulang dengan --model-dir menunjuk folder hasil fine-tuning yang berisi config/tokenizer/model weights."
    else:
        result = generator.generate_from_file(args.input, max_candidates_per_chunk=args.max_candidates)

    payload = json.dumps(_json_safe(result), ensure_ascii=False, indent=2)
    if args.output:
        output_path = Path(args.output)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(payload, encoding="utf-8")
        print(f"JSON kuis tersimpan: {output_path}")
        print(f"Jumlah kuis: {result.get('quiz_count', 0)} | Mode: {result.get('generation_mode')}")
    else:
        print(payload)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
