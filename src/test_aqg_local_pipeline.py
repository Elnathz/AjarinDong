import sys
import unittest
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import aqg_local_pipeline as aqg


class TestAQGPipeline(unittest.TestCase):
    def test_mark_answer_preserves_suffix(self):
        text = "awal target akhir"
        self.assertEqual(aqg.mark_answer(text, "target", 5), "awal <answer>target</answer> akhir")

    def test_formula_marker_excludes_answer_span(self):
        text = "Rumus v = s/t. Hasilnya 12 m/s."
        target = "12 m/s"
        start = text.index(target)
        marked = aqg.mark_answer_and_formulas(text, target, start)
        self.assertIn("<formula>v = s/t</formula>", marked)
        self.assertIn("<answer>12 m/s</answer>", marked)
        self.assertNotIn("<formula><answer>", marked)

    def test_math_unicode_normalization_and_units(self):
        self.assertEqual(aqg.normalize_math_text("12 m/s²"), "12 m/s^2")
        self.assertEqual(aqg.normalize_math_text("40 cm²"), "40 cm^2")
        self.assertEqual(aqg.answer_type("12 m/s^2"), "number")
        self.assertEqual(aqg.answer_type("40 cm^2"), "number")
        self.assertEqual(aqg.answer_type("2024"), "date")
        self.assertEqual(aqg.answer_type("v = s / t"), "formula")

    def test_slide_style_keeps_answer(self):
        context = "Kecepatan dihitung dari jarak dan waktu. Rumusnya v = s/t."
        styled = aqg.convert_to_slide_style(context, "v = s/t")
        self.assertIn("Topik:", styled)
        self.assertIn("Poin Penting:", styled)
        self.assertIn("v = s/t", styled)

    def test_candidate_extractor_gets_formula_and_number(self):
        context = "Kecepatan rata-rata dihitung menggunakan v = s/t. Jika s = 120 m dan t = 10 s, hasilnya 12 m/s."
        rows = aqg.CandidateAnswerExtractor().extract(context, lang="id", max_candidates=5)
        answers = {(row["answer"], row["answer_type"]) for row in rows}
        self.assertTrue(any(kind == "formula" for _, kind in answers))
        self.assertTrue(any(kind == "number" and "m/s" in answer for answer, kind in answers))

    def test_numeric_distractors_keep_same_unit(self):
        distractors = aqg.generate_distractors("12 m/s", "number", "", [])
        self.assertEqual(len(distractors), 3)
        self.assertTrue(all(x.endswith("m/s") for x in distractors))

    def test_formula_distractors(self):
        distractors = aqg.generate_distractors("c^2 = a^2 + b^2", "formula", "", [])
        self.assertEqual(len(distractors), 3)
        self.assertEqual(len(set(distractors)), 3)

    def test_demo_runs_without_model(self):
        generator = aqg.QuizGenerator(load_model=False)
        result = generator.generate_from_chunks(aqg.build_demo_chunks(), max_candidates_per_chunk=5)
        self.assertEqual(result["input_chunks"], 2)
        self.assertGreater(result["candidate_count"], 0)
        self.assertTrue(all(row["question_source"] == "template_fallback" for row in result["quizzes"]))

    @unittest.skipUnless(aqg.fitz is not None, "PyMuPDF not installed")
    def test_pdf_parser_does_not_drop_single_page(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sample.pdf"
            pdf = aqg.fitz.open()
            page = pdf.new_page()
            page.insert_text((72, 72), "Materi Fisika")
            page.insert_text((72, 100), "Kecepatan menggunakan v = s/t dan menghasilkan 12 m/s.")
            pdf.save(path)
            pdf.close()
            chunks = aqg.DocumentToContextPipeline().process(path)
            self.assertEqual(len(chunks), 1)
            self.assertIn("Kecepatan", chunks[0]["context"])

    @unittest.skipUnless(aqg.Presentation is not None, "python-pptx not installed")
    def test_pptx_parser_stitches_slide_bullets_and_table(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "sample.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Gerak Lurus"
            slide.placeholders[1].text = "Kecepatan dihitung dengan v = s/t.\nJarak 120 m dalam 10 s menghasilkan 12 m/s."
            table = slide.shapes.add_table(2, 2, 100, 100, 5000000, 1000000).table
            table.cell(0, 0).text = "Besaran"
            table.cell(0, 1).text = "Nilai"
            table.cell(1, 0).text = "Kecepatan"
            table.cell(1, 1).text = "12 m/s"
            presentation.save(path)
            chunks = aqg.DocumentToContextPipeline().process(path)
            self.assertTrue(chunks)
            self.assertIn("Gerak Lurus", chunks[0]["context"])
            self.assertIn("Data pada tabel", chunks[0]["context"])

    @unittest.skipUnless(aqg.Document is not None, "python-docx not installed")
    def test_docx_omml_fraction(self):
        import xml.etree.ElementTree as ET
        m = "http://schemas.openxmlformats.org/officeDocument/2006/math"
        xml = ET.fromstring(
            f'<m:oMath xmlns:m="{m}"><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num>'
            f'<m:den><m:r><m:t>2</m:t></m:r></m:den></m:f></m:oMath>'
        )
        self.assertEqual(aqg.DocumentToContextPipeline._omml_to_text(xml), "(1)/(2)")


if __name__ == "__main__":
    unittest.main(verbosity=2)
